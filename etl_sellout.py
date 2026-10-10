#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ETL de Sell Out — ISDIN Colombia.

Lee en streaming el Excel del paso 3 del Sell Out (.xlsm), agrega las ventas y
deja en CARPETA_SALIDA los JSON que consume la web app:

  so_pdv.json          PDV × BU × mes
  so_portafolio.json   SKU × mes (con descripción, familia y BU)
  so_indice.json       PDV -> fragmento de detalle, descripción y totales
  so_detalle_NN.json   PDV × SKU × mes, repartido en N_FRAGMENTOS archivos
  so_sku_NN.json       SKU × PDV × mes (índice invertido, rotación por producto),
                       repartido en N_FRAGMENTOS_SKU archivos
  so_sku_indice.json   SKU -> fragmento so_sku_NN, más la lista de meses
  so_manifiesto.json   metadatos de la corrida (se escribe al final)

Las series usan tripletas compactas [índiceMes, unidades, importe], donde
índiceMes es la posición del mes en la lista "meses" del mismo archivo.

Opcionalmente, FILTROS deja pasar solo ciertas filas (p. ej. Affiliate = Colombia)
y DIMENSION_PRODUCTOS completa bu / prod_desc / familia cruzando el SKU con una
tabla de productos de otra hoja del mismo libro.

Uso:
  python etl_sellout.py --inspect       revisa hojas, encabezados y mapeo
  python etl_sellout.py                 corrida completa
  python etl_sellout.py --limite 5000   solo las primeras 5000 filas (pruebas)
  python etl_sellout.py --solo-s3       solo los archivos s3_* de Ventas 3.0 (varios años)
  python etl_sellout.py --sin-s3        solo los archivos so_* del mapa

La configuración vive en config.py (copia de config.example.py).
"""

import argparse
import itertools
import json
import os
import re
import sys
import time
import unicodedata
import zlib
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

CARPETA_SCRIPT = Path(__file__).resolve().parent

CAMPOS_NUMERO = ("units", "amount")
SIN_BU = "SIN_BU"
SIN_SKU = "SIN_SKU"
CADA_N_FILAS = 200_000  # frecuencia del mensaje de progreso

EPOCA_EXCEL = datetime(1899, 12, 30)
RE_YMD = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:$|[ T])")
RE_DMY = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})(?:$|[ T])")


# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

def cargar_config():
    if not (CARPETA_SCRIPT / "config.py").exists():
        print(
            f"ERROR: no se encontró config.py en {CARPETA_SCRIPT}\n"
            "\n"
            "  Este script lee su configuración (ruta del Excel, carpeta de salida\n"
            "  y mapeo de columnas) desde config.py, que NO se sube a GitHub.\n"
            "\n"
            "  Créalo copiando la plantilla y ajusta las rutas a tu equipo:\n"
            "    Windows:    copy config.example.py config.py\n"
            "    Mac/Linux:  cp config.example.py config.py\n"
            "\n"
            "  Luego revisa el mapeo de columnas con: python etl_sellout.py --inspect",
            file=sys.stderr,
        )
        sys.exit(1)

    sys.path.insert(0, str(CARPETA_SCRIPT))
    import config

    faltan = [n for n in ("RUTA_EXCEL", "CARPETA_SALIDA", "COLUMNAS", "OBLIGATORIOS")
              if not hasattr(config, n)]
    if faltan:
        sys.exit(f"ERROR: a config.py le faltan variables: {', '.join(faltan)}. "
                 "Compáralo con config.example.py.")
    return config


# ---------------------------------------------------------------------------
# Lectura del Excel (calamine con respaldo en openpyxl)
# ---------------------------------------------------------------------------

class LibroExcel:
    """Abre el libro con python-calamine si está instalado; si no, con openpyxl
    en modo read_only. En ambos casos las filas se entregan de una en una."""

    def __init__(self, ruta):
        self.ruta = str(ruta)
        try:
            from python_calamine import CalamineWorkbook
        except ImportError:
            CalamineWorkbook = None

        if CalamineWorkbook is not None:
            self.motor = "python-calamine"
            self._libro = CalamineWorkbook.from_path(self.ruta)
            self.hojas = list(self._libro.sheet_names)
            return

        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: no hay motor para leer Excel. Instala las dependencias con:\n"
                     "  pip install -r requirements.txt")
        self.motor = "openpyxl"
        self._libro = load_workbook(self.ruta, read_only=True, data_only=True)
        self.hojas = list(self._libro.sheetnames)

    def filas(self, hoja):
        """Itera las filas de la hoja; la posición 0 de cada fila es la columna A."""
        if self.motor == "openpyxl":
            yield from self._libro[hoja].iter_rows(values_only=True)
            return

        hoja_c = self._libro.get_sheet_by_name(hoja)
        filas = iter(hoja_c.iter_rows() if hasattr(hoja_c, "iter_rows")
                     else hoja_c.to_python(skip_empty_area=False))
        primera = next(filas, None)
        if primera is None:
            return
        # Si la hoja no empieza en A1, según la versión calamine entrega o no las
        # filas/columnas vacías iniciales. Se rellena lo que falte para que la
        # fila 1 y la columna A sigan siendo las del Excel.
        fila0, col0 = getattr(hoja_c, "start", None) or (0, 0)
        col_fin = (getattr(hoja_c, "end", None) or (0, 0))[1]
        if fila0 and not fila_vacia(primera):
            for _ in range(fila0):
                yield ()
        rellenar = col0 > 0 and len(primera) < col_fin + 1
        relleno = [None] * col0
        for fila in itertools.chain([primera], filas):
            yield relleno + list(fila) if rellenar else fila

    def elegir_hoja(self, nombre):
        if not self.hojas:
            sys.exit("ERROR: el libro no tiene hojas.")
        if nombre is None:
            return self.hojas[0]
        if nombre not in self.hojas:
            sys.exit(f"ERROR: la hoja '{nombre}' no existe. Hojas disponibles: "
                     + ", ".join(repr(h) for h in self.hojas))
        return nombre

    def cerrar(self):
        cerrar = getattr(self._libro, "close", None)
        if callable(cerrar):
            cerrar()

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.cerrar()


def abrir_libro(ruta):
    if not ruta.is_file():
        sys.exit(f"ERROR: no existe el archivo Excel: {ruta}\n"
                 "  Revisa RUTA_EXCEL en config.py.")
    try:
        return LibroExcel(ruta)
    except PermissionError:
        sys.exit(f"ERROR: sin permiso para leer {ruta}. ¿Está abierto y bloqueado en Excel?")


def leer_encabezados(libro, hoja, fila_encabezados):
    """Devuelve (encabezados, iterador) con el iterador ya en la primera fila de datos."""
    filas = iter(libro.filas(hoja))
    for n, fila in enumerate(filas, start=1):
        if n == fila_encabezados:
            return [normalizar_texto(c) for c in fila], filas
    sys.exit(f"ERROR: la hoja '{hoja}' tiene menos de {fila_encabezados} filas; "
             "revisa FILA_ENCABEZADOS en config.py.")


def celda(fila, indice):
    if indice is None or indice >= len(fila):
        return None
    return fila[indice]


def fila_vacia(fila):
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in fila)


# ---------------------------------------------------------------------------
# Mapeo de columnas
# ---------------------------------------------------------------------------

def normalizar_encabezado(texto):
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower()
    texto = re.sub(r"[_\-./]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def coincide_parcial(encabezado, fragmento):
    # Los fragmentos cortos solo valen como palabra completa ("bu" no está en "distribuidor").
    if len(fragmento) <= 3:
        return fragmento in encabezado.split()
    return fragmento in encabezado


def mapear_columnas(encabezados, columnas_cfg):
    """Devuelve {campo: índice de columna 0-based} para los campos encontrados.
    Primero coincidencias exactas para todos los campos, luego parciales."""
    normalizados = [normalizar_encabezado(h) for h in encabezados]
    mapeo, usadas = {}, set()

    for comparar in (str.__eq__, coincide_parcial):
        for campo, candidatos in columnas_cfg.items():
            if campo in mapeo:
                continue
            for candidato in candidatos:
                fragmento = normalizar_encabezado(candidato)
                indice = next((i for i, h in enumerate(normalizados)
                               if h and i not in usadas and comparar(h, fragmento)), None)
                if indice is not None:
                    mapeo[campo] = indice
                    usadas.add(indice)
                    break
    return mapeo


def campos_obligatorios(cfg):
    # fecha y pos_id son imprescindibles para agregar, estén o no en OBLIGATORIOS.
    return list(dict.fromkeys(["fecha", "pos_id", *cfg.OBLIGATORIOS]))


def letra_columna(numero):
    letras = ""
    while numero:
        numero, resto = divmod(numero - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


# ---------------------------------------------------------------------------
# Normalización de valores
# ---------------------------------------------------------------------------

def normalizar_texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # códigos numéricos: 1234.0 -> "1234"
    return re.sub(r"\s+", " ", str(valor)).strip()


def normalizar_codigo(valor):
    """POS_ID / SKU: mayúsculas y sin ningún espacio."""
    return re.sub(r"\s+", "", normalizar_texto(valor)).upper()


def normalizar_pos(valor):
    """POS_ID: como normalizar_codigo y además unifica el separador ":" con "_"
    ("111:S566" y "111_S566" son el mismo PDV)."""
    return normalizar_codigo(valor).replace(":", "_")


def normalizar_numero(valor):
    """Número (float) o None. Acepta coma decimal y separadores de miles:
    "1.234.567,89", "1,234,567.89", "12,5", "1.234.567", "$ 1.500", "(1.500)".
    Con un solo separador y sin el otro, la coma se toma como decimal y el
    punto también; si el separador se repite, se toma como de miles."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return None if valor != valor else float(valor)  # descarta NaN

    texto = str(valor).strip().replace("$", "").replace("\u00a0", "").replace(" ", "")
    if not texto:
        return None
    negativo = texto.startswith("(") and texto.endswith(")")
    if negativo:
        texto = texto[1:-1]

    if "," in texto and "." in texto:
        # El separador que aparece de último es el decimal.
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", "") if texto.count(",") > 1 else texto.replace(",", ".")
    elif texto.count(".") > 1:
        texto = texto.replace(".", "")

    try:
        numero = float(texto)
    except ValueError:
        return None
    return -numero if negativo else numero


def _anio_mes(anio, mes, dia):
    try:
        date(anio, mes, dia)
    except ValueError:
        return None
    return anio * 100 + mes


def _serial_excel(serial):
    # Rango razonable de seriales: 1954-10-03 a 2201-03-02.
    if not 20_000 <= serial <= 110_000:
        return None
    fecha = EPOCA_EXCEL + timedelta(days=serial)
    return fecha.year * 100 + fecha.month


def normalizar_fecha(valor):
    """Mes de la fecha como entero AAAAMM, o None si no se reconoce.
    Acepta datetime/date, serial de Excel y texto d/m/yyyy o yyyy-mm-dd."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, date):  # incluye datetime
        return valor.year * 100 + valor.month
    if isinstance(valor, (int, float)):
        return _serial_excel(valor)

    texto = str(valor).strip()
    if not texto:
        return None
    m = RE_YMD.match(texto)
    if m:
        return _anio_mes(int(m[1]), int(m[2]), int(m[3]))
    m = RE_DMY.match(texto)
    if m:
        return _anio_mes(int(m[3]), int(m[2]), int(m[1]))
    try:
        return _serial_excel(float(texto))
    except ValueError:
        return None


def texto_mes(anio_mes):
    return f"{anio_mes // 100:04d}-{anio_mes % 100:02d}"


def normalizar_campo(campo, valor):
    if campo == "fecha":
        anio_mes = normalizar_fecha(valor)
        return texto_mes(anio_mes) if anio_mes else None
    if campo == "pos_id":
        return normalizar_pos(valor)
    if campo == "sku":
        return normalizar_codigo(valor)
    if campo in CAMPOS_NUMERO:
        return normalizar_numero(valor)
    if campo == "bu":
        return normalizar_texto(valor).upper()
    return normalizar_texto(valor)


# ---------------------------------------------------------------------------
# Filtros de filas y dimensión de productos
# ---------------------------------------------------------------------------

CAMPOS_DIMENSION = ("bu", "prod_desc", "familia")


def buscar_columna(encabezados, candidatos):
    """Índice 0-based de la columna que coincide con algún candidato, o None."""
    return mapear_columnas(encabezados, {"columna": candidatos}).get("columna")


def preparar_filtros(cfg, encabezados):
    """Devuelve (filtros, columnas no encontradas). Cada filtro es una tupla
    (índice, encabezado, valores aceptados normalizados, valores tal cual en config)."""
    filtros, faltan = [], []
    for columna, aceptados in (getattr(cfg, "FILTROS", None) or {}).items():
        indice = buscar_columna(encabezados, [columna])
        if indice is None:
            faltan.append(columna)
        else:
            filtros.append((indice, encabezados[indice],
                            {normalizar_encabezado(v) for v in aceptados}, list(aceptados)))
    return filtros, faltan


def pasa_filtros(fila, filtros, cache):
    """True si la fila cumple todos los filtros (sin distinguir mayúsculas ni tildes).
    `cache` guarda el valor normalizado de cada celda ya vista."""
    for indice, _, aceptados, _ in filtros:
        valor = celda(fila, indice)
        clave = cache.get(valor)
        if clave is None:
            if len(cache) > 10_000:
                cache.clear()
            clave = cache[valor] = normalizar_encabezado(normalizar_texto(valor))
        if clave not in aceptados:
            return False
    return True


def cargar_dimension_productos(cfg, libro):
    """Lee la tabla de DIMENSION_PRODUCTOS. Devuelve (productos, info), donde
    productos es {sku: {campo: valor}}; ({}, None) si está desactivada."""
    dim = getattr(cfg, "DIMENSION_PRODUCTOS", None)
    if not dim:
        return {}, None
    sobran = set(dim["campos"]) - set(CAMPOS_DIMENSION)
    if sobran:
        sys.exit(f"ERROR: DIMENSION_PRODUCTOS solo puede completar {', '.join(CAMPOS_DIMENSION)}; "
                 f"sobran: {', '.join(sorted(sobran))}.")

    hoja = libro.elegir_hoja(dim["hoja"])
    encabezados, filas = leer_encabezados(libro, hoja, int(dim.get("fila_encabezados", 1)))
    columnas = mapear_columnas(encabezados, {"clave": dim["clave"], **dim["campos"]})
    faltan = [c for c in ("clave", *dim["campos"]) if c not in columnas]
    if faltan:
        sys.exit(f"ERROR: DIMENSION_PRODUCTOS: en la hoja '{hoja}' no se encontraron columnas "
                 f"para: {', '.join(faltan)}.\n  Encabezados de la hoja: {encabezados}")

    i_clave = columnas.pop("clave")
    productos, repetidos = {}, set()
    for fila in filas:
        sku = normalizar_codigo(celda(fila, i_clave))
        if not sku:
            continue
        if sku in productos:
            repetidos.add(sku)  # gana la primera fila
            continue
        productos[sku] = {campo: normalizar_campo(campo, celda(fila, i))
                          for campo, i in columnas.items()}
    info = {
        "hoja": hoja,
        "clave": encabezados[i_clave],
        "columnas": {campo: encabezados[i] for campo, i in columnas.items()},
        "productos": len(productos),
        "repetidos": sorted(repetidos),
    }
    return productos, info


def completar_con_dimension(valores, productos):
    """Rellena los campos vacíos de la fila con los de la dimensión.
    Devuelve False si el SKU no está en la dimensión."""
    datos = productos.get(valores.get("sku"))
    if datos is None:
        return False
    for campo, valor in datos.items():
        if not valores.get(campo):
            valores[campo] = valor
    return True


# ---------------------------------------------------------------------------
# Modo --inspect
# ---------------------------------------------------------------------------

def inspeccionar(cfg):
    ruta = Path(cfg.RUTA_EXCEL)
    fila_enc = int(getattr(cfg, "FILA_ENCABEZADOS", 1))
    obligatorios = campos_obligatorios(cfg)

    with abrir_libro(ruta) as libro:
        hoja = libro.elegir_hoja(getattr(cfg, "NOMBRE_HOJA", None))
        print(f"Archivo: {ruta}")
        print(f"Motor:   {libro.motor}")
        print(f"\nHojas ({len(libro.hojas)}):")
        for nombre in libro.hojas:
            print(f"  {'->' if nombre == hoja else '  '} {nombre}")

        productos, info_dim = cargar_dimension_productos(cfg, libro)
        encabezados, filas = leer_encabezados(libro, hoja, fila_enc)
        print(f"\nEncabezados de '{hoja}' (fila {fila_enc}):")
        for numero, encabezado in enumerate(encabezados, start=1):
            if encabezado:
                print(f"  {numero:>4}  {letra_columna(numero):>3}  {encabezado}")

        mapeo = mapear_columnas(encabezados, cfg.COLUMNAS)
        print("\nMapeo automático de columnas:")
        for campo in cfg.COLUMNAS:
            marca = "*" if campo in obligatorios else " "
            if campo in mapeo:
                i = mapeo[campo]
                destino = f"columna {i + 1} ({letra_columna(i + 1)}) '{encabezados[i]}'"
            else:
                destino = "NO ENCONTRADA"
            print(f"  {marca} {campo:<10} -> {destino}")
        print("  (* = obligatorio)")

        faltan = [c for c in obligatorios if c not in mapeo]
        if faltan:
            print(f"\nATENCIÓN: faltan campos obligatorios: {', '.join(faltan)}. "
                  "Agrega fragmentos a COLUMNAS en config.py.")

        filtros, faltan_filtros = preparar_filtros(cfg, encabezados)
        print("\nFiltros de filas:")
        if not filtros and not faltan_filtros:
            print("  (sin filtros: se procesan todas las filas)")
        for indice, encabezado, _, aceptados in filtros:
            print(f"  columna {indice + 1} ({letra_columna(indice + 1)}) '{encabezado}' "
                  f"debe ser: {', '.join(map(str, aceptados))}")
        for columna in faltan_filtros:
            print(f"  ATENCIÓN: no se encontró la columna del filtro '{columna}'. Revisa FILTROS.")

        print("\nDimensión de productos:")
        if info_dim is None:
            print("  (desactivada)")
        else:
            print(f"  hoja '{info_dim['hoja']}': {info_dim['productos']:,} productos, "
                  f"cruce del sku con la columna '{info_dim['clave']}'")
            for campo, encabezado in info_dim["columnas"].items():
                print(f"    {campo:<10} <- '{encabezado}' (solo si la venta lo trae vacío)")
            if info_dim["repetidos"]:
                print(f"  {len(info_dim['repetidos'])} códigos repetidos (se usa la primera fila): "
                      + ", ".join(info_dim["repetidos"]))

        print("\nPrimeras 3 filas de datos" + (" que pasan los filtros:" if filtros else ":"))
        mostradas, cache = 0, {}
        for n, fila in enumerate(filas, start=fila_enc + 1):
            if fila_vacia(fila) or not pasa_filtros(fila, filtros, cache):
                continue
            mostradas += 1
            print(f"\n  Fila {n}:")
            for i, valor in enumerate(fila):
                if valor is not None and str(valor).strip():
                    nombre = encabezados[i] if i < len(encabezados) else ""
                    print(f"    {i + 1:>4}  {nombre[:30]:<30}  {valor!r}")
            valores = {campo: normalizar_campo(campo, celda(fila, mapeo.get(campo)))
                       for campo in cfg.COLUMNAS}
            originales = dict(valores)
            cruzo = completar_con_dimension(valores, productos) if productos else None
            print("    Normalizado:")
            for campo in cfg.COLUMNAS:
                completado = valores[campo] != originales[campo]
                if campo in mapeo or completado:
                    origen = f"   <- {info_dim['hoja']}" if completado else ""
                    print(f"      {campo:<10} = {valores[campo]!r}{origen}")
            if cruzo is False:
                print(f"      (el sku no está en '{info_dim['hoja']}')")
            if mostradas == 3:
                break
        if mostradas == 0:
            print("  (no hay filas de datos" + (" que pasen los filtros)" if filtros else ")"))


# ---------------------------------------------------------------------------
# Modo normal: agregación y escritura de JSON
# ---------------------------------------------------------------------------

def redondear(numero):
    valor = round(numero, 2)
    return int(valor) if valor.is_integer() else valor


def fragmento_de(pos_id, n_fragmentos):
    # crc32 es estable entre corridas (hash() de Python no lo es).
    return zlib.crc32(pos_id.encode("utf-8")) % n_fragmentos


def escribir_json(ruta, datos, intentos=5, ascii=False):
    """Escritura atómica (archivo .tmp + reemplazo) para que Drive nunca
    sincronice un JSON a medio escribir. Reintenta si Drive bloquea el archivo.
    ascii=True escapa tildes y Ñ (\\uXXXX): la web app cachea por trozos de 90 KB
    solo si el texto es ASCII (con tildes, de 30 KB, y los archivos grandes ya no caben)."""
    temporal = ruta.with_name(ruta.name + ".tmp")
    with open(temporal, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=ascii, separators=(",", ":"))
    for intento in range(intentos):
        try:
            os.replace(temporal, ruta)
            break
        except PermissionError:
            if intento == intentos - 1:
                raise
            time.sleep(1)
    return ruta.stat().st_size


def ejecutar(cfg, limite):
    t0 = time.perf_counter()
    ruta = Path(cfg.RUTA_EXCEL)
    salida = Path(cfg.CARPETA_SALIDA)
    fila_enc = int(getattr(cfg, "FILA_ENCABEZADOS", 1))
    n_fragmentos = int(getattr(cfg, "N_FRAGMENTOS", 16))
    n_fragmentos_sku = int(getattr(cfg, "N_FRAGMENTOS_SKU", 8))
    obligatorios = campos_obligatorios(cfg)

    # Acumuladores [unidades, importe] por mes (clave entera AAAAMM).
    por_pdv_bu = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0.0, 0.0])))
    por_sku = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0]))
    por_pdv_sku = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0.0, 0.0])))
    # Índice invertido SKU × PDV × mes (rotación por producto).
    por_sku_pdv = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0.0, 0.0])))
    desc_pdv = {}
    info_sku = defaultdict(dict)
    conteo = defaultdict(int)

    with abrir_libro(ruta) as libro:
        hoja = libro.elegir_hoja(getattr(cfg, "NOMBRE_HOJA", None))
        motor = libro.motor
        productos, info_dim = cargar_dimension_productos(cfg, libro)
        encabezados, filas = leer_encabezados(libro, hoja, fila_enc)
        mapeo = mapear_columnas(encabezados, cfg.COLUMNAS)

        faltan = [c for c in obligatorios if c not in mapeo]
        if faltan:
            sys.exit(f"ERROR: no se encontraron columnas para: {', '.join(faltan)}.\n"
                     "  Revisa COLUMNAS en config.py con: python etl_sellout.py --inspect")
        filtros, faltan_filtros = preparar_filtros(cfg, encabezados)
        if faltan_filtros:
            sys.exit(f"ERROR: no se encontraron las columnas de FILTROS: {', '.join(faltan_filtros)}.\n"
                     "  Revisa FILTROS en config.py con: python etl_sellout.py --inspect")

        print(f"Leyendo '{hoja}' de {ruta.name} con {motor}"
              + (f" (límite: {limite:,} filas)" if limite else "") + " ...")
        for _, encabezado, _, aceptados in filtros:
            print(f"  Filtro: '{encabezado}' = {', '.join(map(str, aceptados))}")
        if info_dim:
            print(f"  Dimensión: {info_dim['productos']:,} productos de '{info_dim['hoja']}'")

        col = {campo: mapeo.get(campo) for campo in cfg.COLUMNAS}
        opcionales = ("amount", "units", "bu", "sku", "pdv_desc", "prod_desc", "familia")
        exigidos = [c for c in opcionales if c in obligatorios]
        cache_filtros = {}
        skus_sin_dimension = defaultdict(int)

        for fila in filas:
            if fila_vacia(fila):
                conteo["vacias"] += 1
                continue
            if filtros and not pasa_filtros(fila, filtros, cache_filtros):
                conteo["filtradas"] += 1
                continue
            # --limite cuenta solo filas que pasan los filtros.
            if limite is not None and conteo["leidas"] >= limite:
                break
            conteo["leidas"] += 1
            if conteo["leidas"] % CADA_N_FILAS == 0:
                print(f"  {conteo['leidas']:,} filas ({time.perf_counter() - t0:.0f} s)")

            anio_mes = normalizar_fecha(celda(fila, col.get("fecha")))
            if anio_mes is None:
                conteo["descartadas_fecha"] += 1
                continue
            pos_id = normalizar_pos(celda(fila, col.get("pos_id")))
            if not pos_id:
                conteo["descartadas_pos_id"] += 1
                continue

            valores = {
                "amount": normalizar_numero(celda(fila, col.get("amount"))),
                "units": normalizar_numero(celda(fila, col.get("units"))),
                "bu": normalizar_texto(celda(fila, col.get("bu"))).upper(),
                "sku": normalizar_codigo(celda(fila, col.get("sku"))),
                "pdv_desc": normalizar_texto(celda(fila, col.get("pdv_desc"))),
                "prod_desc": normalizar_texto(celda(fila, col.get("prod_desc"))),
                "familia": normalizar_texto(celda(fila, col.get("familia"))),
            }
            if productos and valores["sku"] and not completar_con_dimension(valores, productos):
                skus_sin_dimension[valores["sku"]] += 1
            faltante = next((c for c in exigidos if valores[c] in (None, "")), None)
            if faltante:
                conteo[f"descartadas_{faltante}"] += 1
                continue

            unidades = valores["units"] or 0.0
            importe = valores["amount"] or 0.0
            bu = valores["bu"] or SIN_BU
            sku = valores["sku"] or SIN_SKU

            for acumulado in (por_pdv_bu[pos_id][bu][anio_mes],
                              por_sku[sku][anio_mes],
                              por_pdv_sku[pos_id][sku][anio_mes],
                              por_sku_pdv[sku][pos_id][anio_mes]):
                acumulado[0] += unidades
                acumulado[1] += importe

            if valores["pdv_desc"] and pos_id not in desc_pdv:
                desc_pdv[pos_id] = valores["pdv_desc"]
            info = info_sku[sku]
            if valores["prod_desc"] and "desc" not in info:
                info["desc"] = valores["prod_desc"]
            if valores["familia"] and "familia" not in info:
                info["familia"] = valores["familia"]
            if bu != SIN_BU and "bu" not in info:
                info["bu"] = bu
            conteo["validas"] += 1

    t_lectura = time.perf_counter() - t0
    print(f"Lectura terminada: {conteo['leidas']:,} filas, {conteo['validas']:,} válidas"
          + (f", {conteo['filtradas']:,} excluidas por filtros" if filtros else "")
          + f" ({t_lectura:.1f} s). Escribiendo JSON en {salida} ...")

    # índiceMes = posición en la lista ordenada de meses con datos.
    anios_mes = sorted({am for por_mes in por_sku.values() for am in por_mes})
    meses = [texto_mes(am) for am in anios_mes]
    indice_mes = {am: i for i, am in enumerate(anios_mes)}

    def serie(por_mes):
        return [[indice_mes[am], redondear(u), redondear(a)]
                for am, (u, a) in sorted(por_mes.items())]

    total_mes = defaultdict(lambda: [0.0, 0.0])
    so_pdv, so_indice = {}, {}
    detalle = [{} for _ in range(n_fragmentos)]
    for pos_id in sorted(por_pdv_bu):
        bus = por_pdv_bu[pos_id]
        total_u = total_a = 0.0
        for por_mes in bus.values():
            for am, (u, a) in por_mes.items():
                total_u += u
                total_a += a
                total_mes[am][0] += u
                total_mes[am][1] += a
        fragmento = fragmento_de(pos_id, n_fragmentos)
        descripcion = desc_pdv.get(pos_id, "")
        so_pdv[pos_id] = {"desc": descripcion,
                          "bu": {bu: serie(bus[bu]) for bu in sorted(bus)}}
        so_indice[pos_id] = {"desc": descripcion, "frag": fragmento,
                             "unidades": redondear(total_u), "importe": redondear(total_a)}
        skus = por_pdv_sku[pos_id]
        detalle[fragmento][pos_id] = {sku: serie(skus[sku]) for sku in sorted(skus)}

    so_portafolio = {}
    for sku in sorted(por_sku):
        info = info_sku.get(sku, {})
        so_portafolio[sku] = {"desc": info.get("desc", ""), "familia": info.get("familia", ""),
                              "bu": info.get("bu", ""), "serie": serie(por_sku[sku])}

    salida.mkdir(parents=True, exist_ok=True)
    archivos = []

    def guardar(nombre, datos):
        archivos.append({"nombre": nombre, "bytes": escribir_json(salida / nombre, datos)})

    guardar("so_pdv.json", {"meses": meses, "pdv": so_pdv})
    guardar("so_portafolio.json", {"meses": meses, "sku": so_portafolio})
    guardar("so_indice.json", {"n_fragmentos": n_fragmentos, "pdv": so_indice})
    for n, pdvs in enumerate(detalle):
        guardar(f"so_detalle_{n:02d}.json", {"fragmento": n, "meses": meses, "pdv": pdvs})

    # Índice invertido producto -> PDV: so_sku_NN.json = {sku: {pos_id: serie}}.
    frags_sku = [{} for _ in range(n_fragmentos_sku)]
    indice_sku = {}
    for sku in sorted(por_sku_pdv):
        pdvs = por_sku_pdv[sku]
        fragmento = fragmento_de(sku, n_fragmentos_sku)
        indice_sku[sku] = fragmento
        frags_sku[fragmento][sku] = {pos_id: serie(pdvs[pos_id]) for pos_id in sorted(pdvs)}
    for n, skus in enumerate(frags_sku):
        guardar(f"so_sku_{n:02d}.json", skus)
    guardar("so_sku_indice.json",
            {"fragmentos": n_fragmentos_sku, "meses": meses, "sku": indice_sku})

    total_u = sum(u for u, _ in total_mes.values())
    total_a = sum(a for _, a in total_mes.values())
    manifiesto = {
        "version": 1,
        "generado": datetime.now().isoformat(timespec="seconds"),
        "origen": {"archivo": ruta.name, "hoja": hoja,
                   "fila_encabezados": fila_enc, "motor": motor},
        "limite": limite,
        "formato_serie": ["indiceMes", "unidades", "importe"],
        "meses": meses,
        "n_fragmentos": n_fragmentos,
        "n_fragmentos_sku": n_fragmentos_sku,
        "mapeo_columnas": {
            campo: ({"columna": mapeo[campo] + 1, "encabezado": encabezados[mapeo[campo]]}
                    if campo in mapeo else None)
            for campo in cfg.COLUMNAS
        },
        "filtros": {encabezado: aceptados for _, encabezado, _, aceptados in filtros},
        "dimension_productos": (
            dict(info_dim, skus_sin_cruce=dict(sorted(skus_sin_dimension.items())))
            if info_dim else None
        ),
        "filas": dict(sorted(conteo.items())),
        "conteos": {"pdv": len(so_pdv), "sku": len(so_portafolio),
                    "bu": len({bu for bus in por_pdv_bu.values() for bu in bus})},
        "totales": {"unidades": redondear(total_u), "importe": redondear(total_a)},
        "totales_mes": serie(total_mes),
        "archivos": archivos,
        "duracion_seg": round(time.perf_counter() - t0, 1),
    }
    # El manifiesto va de último: si está actualizado, el resto de archivos también.
    guardar("so_manifiesto.json", manifiesto)

    print(f"Listo en {time.perf_counter() - t0:.1f} s: {len(meses)} meses"
          + (f" ({meses[0]} a {meses[-1]})" if meses else "")
          + f", {len(so_pdv):,} PDV, {len(so_portafolio):,} SKU.")
    print(f"  Totales: {total_u:,.2f} unidades, {total_a:,.2f} de importe.")
    descartadas = {k[len("descartadas_"):]: v for k, v in conteo.items()
                   if k.startswith("descartadas_")}
    if descartadas:
        print("  Filas descartadas por campo: "
              + ", ".join(f"{k}={v:,}" for k, v in sorted(descartadas.items())))
    if skus_sin_dimension:
        print(f"  {len(skus_sin_dimension)} SKU sin cruce en '{info_dim['hoja']}': "
              + ", ".join(f"{s} ({n:,} filas)" for s, n in sorted(skus_sin_dimension.items())))
    print(f"  {len(archivos)} archivos JSON escritos en {salida}")


# ---------------------------------------------------------------------------
# Ventas 3.0: cubos de sell-out de VARIOS años, con cliente y KAM
# ---------------------------------------------------------------------------
# La pestaña "Ventas 3.0" de la web app (JsVentas3.html) replica el visual de
# crecimiento del Power BI de sell-out (PY · AC · ΔPY · ΔPY% · var. vs mes
# pasado) por Cliente, BU, Producto o Punto de venta. Necesita el año anterior y
# el cliente de cada fila, que los archivos del mapa no traen; por eso lee la
# hoja 'Final' de cada "3. Affiliate_Master so" listado en FUENTES_S3 (uno por
# año). Ahí el cliente es Origin + Sold To ID (= SAP ID), y el punto de venta es
# el POS_ID (SAP ID + "_" + código interno del cliente, la misma clave del mapa y
# del CRM). El libro de 2025 no trae POS_ID: se toma el que ese SF_ID tiene en 2026
# (≈94 % del importe de 2025) y, si no, las reglas de resolver_pos_id. No toca
# ninguno de los archivos so_*.json del mapa.
#
# Archivos (prefijo s3_, JSON ASCII para que la web app los pueda cachear):
#   s3_cubo.json      catálogos + cliente×producto×mes (cs) + PDV×mes (pm)
#   s3_pb_NN.json     PDV × mes de UN BU (NN = índice en "bus"; dimensión PDV filtrada por BU)
#   s3_pdv_NN.json    PDV × producto × mes, por PDV   (filtro de punto de venta)
#   s3_sku_NN.json    PDV × producto × mes, por producto (PDV filtrado por producto)
# Todas las series van como [índices..., mes, unidades, importe] y todos los
# archivos llevan el mismo "generado": la web app no mezcla corridas distintas.

COLUMNAS_S3_DEF = {
    "fecha": ["fecha"],
    "cliente_id": ["sold to id"],
    "cliente": ["origin"],
    "pdv": ["sf id"],
    "pos_id": ["pos id"],
    "pdv_desc": ["isdin pdv desc"],
    "sku": ["ean isdin"],
    "prod_desc": ["ean description"],
    "units": ["units"],
    "amount": ["amount"],
    "kam": ["kam"],
    "canal": ["canal"],
    "subcanal": ["sub canal"],
    "ciudad": ["ciudad cliente"],
    "departamento": ["departamento cliente"],
}


# Código interno del PDV al inicio de su descripción: "883_CRUZ_VERDE…", "CC100 MEDIPIEL…",
# "1094 - FARMATODO…", "ZC_BELLA PIEL…" (en 2026 coincide con el POS_ID en el 94 % de las filas).
RE_COD_PDV = re.compile(r"^\s*([A-Za-z]{0,3}\d+[A-Za-z]?|[A-Za-z]{2})\s*[-_ ]")
_cod_pdv_cache = {}


def codigo_pdv_desc(desc):
    if desc not in _cod_pdv_cache:
        m = RE_COD_PDV.match(desc or "")
        _cod_pdv_cache[desc] = m.group(1) if m else ""
    return _cod_pdv_cache[desc]


def resolver_pos_id(tok, sf_pos):
    """POS_ID (= SAP ID del cliente + "_" + código interno del PDV) de una clave provisional
    (pos, sf, cliente, código de la descripción). El libro de 2025 no trae POS_ID, así que:
      1. el POS_ID de la fila, si lo trae (2026);
      2. el POS_ID que ese mismo SF_ID tiene en las filas que sí lo traen (el más frecuente);
      3. el SF_ID, si ya viene como "SAPID_código";
      4. SAP ID + el código al inicio de la descripción;
      5. si nada de eso, el SF_ID tal cual.
    Devuelve (clave, regla)."""
    pos, sf, cli, cod = tok
    if pos:
        return pos, "pos_id"
    if sf in sf_pos:
        return max(sf_pos[sf].items(), key=lambda kv: kv[1])[0], "sf_id_en_otro_anio"
    sfn = normalizar_pos(sf)
    if "_" in sfn and sfn.startswith(cli + "_"):
        return sfn, "sf_id_compuesto"
    if cod:
        return normalizar_pos(f"{cli}_{cod}"), "codigo_descripcion"
    return sf, "sin_pos_id"


def _ultimo(info, clave, anio_mes, valor):
    """Guarda en info[clave] el valor del mes más reciente (empates: el primero)."""
    if not valor:
        return
    previo = info.get(clave)
    if previo is None or anio_mes > previo[0]:
        info[clave] = (anio_mes, valor)


def ejecutar_s3(cfg, limite=None):
    fuentes = getattr(cfg, "FUENTES_S3", None)
    if not fuentes:
        print("Ventas 3.0: FUENTES_S3 no está en config.py; no se generan los archivos s3_*.")
        return
    t0 = time.perf_counter()
    salida = Path(cfg.CARPETA_SALIDA)
    columnas_cfg = getattr(cfg, "COLUMNAS_S3", None) or COLUMNAS_S3_DEF
    n_frag_pdv = int(getattr(cfg, "N_FRAGMENTOS_S3_PDV", 32))
    n_frag_sku = int(getattr(cfg, "N_FRAGMENTOS_S3_SKU", 16))

    cs = defaultdict(lambda: [0.0, 0.0])        # (cliente, sku, mes)
    # PDV: durante la lectura, una clave PROVISIONAL tok = (pos, sf, cliente, código de la
    # descripción); al final cada tok se resuelve a su POS_ID (resolver_pos_id) y se juntan.
    pm = defaultdict(lambda: [0.0, 0.0])        # (tok, mes)
    ps = defaultdict(lambda: [0.0, 0.0])        # (tok, sku, mes)
    sf_pos = defaultdict(lambda: defaultdict(int))   # SF_ID -> {POS_ID: filas}, de las filas que traen los dos
    filas_tok = defaultdict(int)
    info_cli, info_pdv, info_sku = defaultdict(dict), defaultdict(dict), defaultdict(dict)
    kam_mes = defaultdict(lambda: defaultdict(int))   # (cliente, mes) -> {kam: filas}
    productos = {}
    conteo = defaultdict(int)
    origen = []

    for fuente in fuentes:
        ruta = Path(fuente["ruta"])
        with abrir_libro(ruta) as libro:
            hoja = libro.elegir_hoja(fuente.get("hoja", "Final"))
            dim, info_dim = cargar_dimension_productos(cfg, libro) if getattr(cfg, "DIMENSION_PRODUCTOS", None) and \
                cfg.DIMENSION_PRODUCTOS["hoja"] in libro.hojas else ({}, None)
            productos.update(dim)   # la fuente más reciente (última de la lista) manda
            encabezados, filas = leer_encabezados(libro, hoja, int(fuente.get("fila_encabezados", 1)))
            mapeo = mapear_columnas(encabezados, columnas_cfg)
            faltan = [c for c in ("fecha", "sku", "units", "amount") if c not in mapeo]
            if "cliente_id" not in mapeo and "cliente" not in mapeo:
                faltan.append("cliente_id/cliente")
            if "pdv" not in mapeo and "pos_id" not in mapeo:
                faltan.append("pdv/pos_id")
            if faltan:
                sys.exit(f"ERROR Ventas 3.0: en '{ruta.name}' › '{hoja}' no encontré: {', '.join(faltan)}.\n"
                         f"  Encabezados: {encabezados}")
            filtros, faltan_filtros = preparar_filtros(cfg, encabezados)
            if faltan_filtros:
                sys.exit(f"ERROR Ventas 3.0: en '{ruta.name}' no están las columnas de FILTROS: "
                         + ", ".join(faltan_filtros))
            print(f"Ventas 3.0: leyendo '{hoja}' de {ruta.name} con {libro.motor} ...")
            col = {campo: mapeo.get(campo) for campo in columnas_cfg}
            cache_filtros, leidas = {}, 0
            for fila in filas:
                if fila_vacia(fila):
                    continue
                if filtros and not pasa_filtros(fila, filtros, cache_filtros):
                    conteo["filtradas"] += 1
                    continue
                if limite is not None and leidas >= limite:
                    break
                leidas += 1
                am = normalizar_fecha(celda(fila, col["fecha"]))
                if am is None:
                    conteo["descartadas_fecha"] += 1
                    continue
                nombre_cli = normalizar_texto(celda(fila, col.get("cliente")))
                cli = normalizar_codigo(celda(fila, col.get("cliente_id"))) or normalizar_encabezado(nombre_cli).upper()
                sf = normalizar_texto(celda(fila, col.get("pdv")))
                pos = normalizar_pos(celda(fila, col.get("pos_id")))
                desc_pdv = normalizar_texto(celda(fila, col.get("pdv_desc")))
                sku = normalizar_codigo(celda(fila, col["sku"]))
                if not cli or not (sf or pos) or not sku:
                    conteo["descartadas_" + ("cliente" if not cli else ("pdv" if not (sf or pos) else "sku"))] += 1
                    continue
                if pos and sf:
                    sf_pos[sf][pos] += 1
                tok = (pos, sf, cli, codigo_pdv_desc(desc_pdv))
                filas_tok[tok] += 1
                u = normalizar_numero(celda(fila, col["units"])) or 0.0
                a = normalizar_numero(celda(fila, col["amount"])) or 0.0
                for acumulado in (cs[(cli, sku, am)], pm[(tok, am)], ps[(tok, sku, am)]):
                    acumulado[0] += u
                    acumulado[1] += a
                ic, ip, isk = info_cli[cli], info_pdv[tok], info_sku[sku]
                _ultimo(ic, "nombre", am, nombre_cli)
                kam = normalizar_texto(celda(fila, col.get("kam"))).upper()
                if kam:
                    kam_mes[(cli, am)][kam] += 1
                _ultimo(ip, "desc", am, desc_pdv)
                _ultimo(ip, "cli", am, cli)
                _ultimo(ip, "ciudad", am, normalizar_texto(celda(fila, col.get("ciudad"))).upper())
                _ultimo(ip, "depto", am, normalizar_texto(celda(fila, col.get("departamento"))).upper())
                _ultimo(ip, "canal", am, normalizar_texto(celda(fila, col.get("canal"))).upper())
                _ultimo(ip, "subcanal", am, normalizar_texto(celda(fila, col.get("subcanal"))).upper())
                _ultimo(ip, "pos", am, normalizar_pos(celda(fila, col.get("pos_id"))))
                _ultimo(isk, "desc", am, normalizar_texto(celda(fila, col.get("prod_desc"))))
                conteo["validas"] += 1
            origen.append({"archivo": ruta.name, "hoja": hoja, "filas_validas": leidas,
                           "columnas": {c: encabezados[i] for c, i in mapeo.items()}})

    # PDV = POS_ID: cada clave provisional se resuelve y se juntan las que caen en el mismo
    # POS_ID (descripción, ciudad, etc. = las del mes más reciente).
    clave_pdv, reglas_pdv = {}, defaultdict(int)
    for tok in info_pdv:
        clave_pdv[tok], regla = resolver_pos_id(tok, sf_pos)
        reglas_pdv[regla] += filas_tok[tok]
    pm_tok, ps_tok, info_tok = pm, ps, info_pdv
    pm, ps, info_pdv = defaultdict(lambda: [0.0, 0.0]), defaultdict(lambda: [0.0, 0.0]), defaultdict(dict)
    for (tok, am), v in pm_tok.items():
        x = pm[(clave_pdv[tok], am)]
        x[0] += v[0]
        x[1] += v[1]
    for (tok, s, am), v in ps_tok.items():
        x = ps[(clave_pdv[tok], s, am)]
        x[0] += v[0]
        x[1] += v[1]
    for tok, inf in info_tok.items():
        destino = info_pdv[clave_pdv[tok]]
        for k, (am, valor) in inf.items():
            if k not in destino or am > destino[k][0]:
                destino[k] = (am, valor)
    del pm_tok, ps_tok, info_tok

    # KAM del cliente = el que más filas tiene en su mes más reciente.
    kam_cli = {}
    for (cli, am), kams in kam_mes.items():
        if cli not in kam_cli or am > kam_cli[cli][0]:
            kam_cli[cli] = (am, max(kams.items(), key=lambda kv: kv[1])[0])

    def val(info, clave):
        return info[clave][1] if clave in info else ""

    anios_mes = sorted({k[2] for k in cs})
    meses = [texto_mes(am) for am in anios_mes]
    i_mes = {am: i for i, am in enumerate(anios_mes)}
    clientes = sorted(info_cli, key=lambda c: (val(info_cli[c], "nombre") or c))
    i_cli = {c: i for i, c in enumerate(clientes)}
    skus = sorted(info_sku, key=lambda s: (val(info_sku[s], "desc") or s))
    i_sku = {s: i for i, s in enumerate(skus)}
    pdvs = sorted(info_pdv)
    i_pdv = {p: i for i, p in enumerate(pdvs)}
    # PDV × BU sale del detalle con el BU FINAL de cada producto (el de la fuente
    # más reciente), igual que los productos del cubo.
    def bu_de(sku):
        return ((productos.get(sku) or {}).get("bu") or SIN_BU).upper()
    pb = defaultdict(lambda: [0.0, 0.0])
    for (p, s, am), v in ps.items():
        x = pb[(p, bu_de(s), am)]
        x[0] += v[0]
        x[1] += v[1]
    bus = sorted({k[1] for k in pb})
    i_bu = {b: i for i, b in enumerate(bus)}
    generado = datetime.now().isoformat(timespec="seconds")

    def r(x):
        return redondear(x)

    cubo = {
        "v": 1, "generado": generado, "meses": meses,
        "clientes": [[c, val(info_cli[c], "nombre") or c, kam_cli.get(c, (0, "SIN KAM"))[1]] for c in clientes],
        "skus": [[s, val(info_sku[s], "desc") or s, bu_de(s), ""] for s in skus],
        "pdv": [[p, val(info_pdv[p], "desc") or p, i_cli.get(val(info_pdv[p], "cli"), -1), val(info_pdv[p], "ciudad"),
                 val(info_pdv[p], "depto"), val(info_pdv[p], "canal"), val(info_pdv[p], "subcanal"),
                 val(info_pdv[p], "pos")] for p in pdvs],
        "cs": [[i_cli[c], i_sku[s], i_mes[am], r(v[0]), r(v[1])] for (c, s, am), v in sorted(cs.items())],
        "pm": [[i_pdv[p], i_mes[am], r(v[0]), r(v[1])] for (p, am), v in sorted(pm.items())],
        "bus": bus,
        "fragmentos": {"pdv": n_frag_pdv, "sku": n_frag_sku},
        "meta": {"fuentes": origen, "filas": dict(sorted(conteo.items())),
                 "filtros": dict(getattr(cfg, "FILTROS", None) or {}),
                 "pdv_pos_id": dict(sorted(reglas_pdv.items()))},   # filas por regla de resolver_pos_id
    }
    frag_pdv = [defaultdict(list) for _ in range(n_frag_pdv)]
    frag_sku = [defaultdict(list) for _ in range(n_frag_sku)]
    for (p, s, am), v in sorted(ps.items()):
        ip, isk, im = i_pdv[p], i_sku[s], i_mes[am]
        frag_pdv[ip % n_frag_pdv][ip].append([isk, im, r(v[0]), r(v[1])])
        frag_sku[isk % n_frag_sku][isk].append([ip, im, r(v[0]), r(v[1])])

    salida.mkdir(parents=True, exist_ok=True)
    escritos = []

    def guardar(nombre, datos):
        escritos.append((nombre, escribir_json(salida / nombre, datos, ascii=True)))

    # Un archivo por BU (todos juntos pasan de 5 MB y la web app ya no los podría cachear).
    por_bu = [[] for _ in bus]
    for (p, b, am), v in sorted(pb.items()):
        por_bu[i_bu[b]].append([i_pdv[p], i_mes[am], r(v[0]), r(v[1])])
    for n, filas_bu in enumerate(por_bu):
        guardar(f"s3_pb_{n:02d}.json", {"generado": generado, "bu": bus[n], "pb": filas_bu})
    for n, d in enumerate(frag_pdv):
        guardar(f"s3_pdv_{n:02d}.json", {"generado": generado, "filas": d})
    for n, d in enumerate(frag_sku):
        guardar(f"s3_sku_{n:02d}.json", {"generado": generado, "filas": d})
    guardar("s3_cubo.json", cubo)   # el cubo va de último: si está nuevo, el resto también

    total = sum(b for _, b in escritos)
    print(f"Ventas 3.0 listo en {time.perf_counter() - t0:.1f} s: {len(meses)} meses"
          + (f" ({meses[0]} a {meses[-1]})" if meses else "")
          + f", {len(clientes)} clientes, {len(pdvs):,} PDV, {len(skus)} productos; "
          f"{len(escritos)} archivos s3_*, {total / 1e6:.1f} MB (s3_cubo.json "
          f"{dict(escritos)['s3_cubo.json'] / 1e6:.1f} MB).")
    descartadas = {k: v for k, v in conteo.items() if k.startswith("descartadas_")}
    if descartadas:
        print("  Filas descartadas: " + ", ".join(f"{k[12:]}={v:,}" for k, v in sorted(descartadas.items())))


# ---------------------------------------------------------------------------

def main():
    for flujo in (sys.stdout, sys.stderr):
        if hasattr(flujo, "reconfigure"):
            flujo.reconfigure(errors="replace")

    parser = argparse.ArgumentParser(description="ETL de Sell Out — ISDIN Colombia.")
    parser.add_argument("--inspect", action="store_true",
                        help="lista hojas, encabezados, mapeo de columnas y 3 filas de muestra")
    parser.add_argument("--limite", type=int, metavar="N",
                        help="procesa solo las primeras N filas de datos (pruebas)")
    parser.add_argument("--solo-s3", action="store_true",
                        help="genera solo los archivos s3_* de Ventas 3.0 (no toca los del mapa)")
    parser.add_argument("--sin-s3", action="store_true",
                        help="genera solo los archivos so_* del mapa (como antes)")
    args = parser.parse_args()
    if args.limite is not None and args.limite <= 0:
        parser.error("--limite debe ser un entero positivo")

    cfg = cargar_config()
    if args.inspect:
        inspeccionar(cfg)
        return
    if not args.solo_s3:
        ejecutar(cfg, args.limite)
    if not args.sin_s3:
        ejecutar_s3(cfg, args.limite)


if __name__ == "__main__":
    main()
